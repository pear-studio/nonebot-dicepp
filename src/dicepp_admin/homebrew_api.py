"""群私设管理 API（admin 后台）

操作 data/instances/<instance>/bots/<bot>/group_homebrew/<group>/*.db。
跟 feat/homebrew-data-layer + feat/homebrew-commands 在 DicePP 插件
里实现的私设系统是一对的：

  插件侧 — 由 .hb add/del/list/宏 等指令写入
  admin 侧（本文件） — 让主持人在 WebUI 直接拖 xlsx 上传、点表格 CRUD

跟其他 admin 模块一样：从不 import DicePP plugin code，全部通过文件
+ sqlite3 操作，跟 .hb 命令解耦但共用同一份磁盘上的 .db。
"""
import os
import shutil
import sqlite3
from contextlib import contextmanager
from pathlib import Path
from typing import Dict, Iterator, List, Optional

from dicepp_admin.config import AdminPaths


# 跟 core/data/query_store.py 保持同一份 schema
QUERY_DATA_FIELDS = ["名称", "英文", "来源", "分类", "标签", "内容"]
QUERY_REDIRECT_FIELDS = ["名称", "重定向"]


def _bots_root(instance_id: str) -> Path:
    return AdminPaths.instance_dir(instance_id) / "bots"


def _homebrew_root(instance_id: str, bot_id: str) -> Path:
    return _bots_root(instance_id) / bot_id / "group_homebrew"


def _group_dir(instance_id: str, bot_id: str, group_id: str) -> Path:
    return _homebrew_root(instance_id, bot_id) / group_id


def _safe_group_id(s: str) -> bool:
    return bool(s) and all(c.isdigit() or c in ("_", "-") for c in s)


def _safe_db_filename(name: str) -> Optional[Path]:
    """name 必须是 *.db 文件名（不含路径分隔符）。"""
    if not name or "/" in name or "\\" in name or ".." in name:
        return None
    if not name.lower().endswith(".db"):
        return None
    return Path(name)


@contextmanager
def _open_db(path: Path) -> Iterator[Optional[sqlite3.Connection]]:
    if not path.exists():
        yield None
        return
    conn = sqlite3.connect(str(path))
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()


# ─── 列表 ────────────────────────────────────────────────────────────────

def list_groups_with_homebrew(instance_id: str, bot_id: str) -> List[Dict]:
    """列出该实例下指定 bot 有哪些群配了私设。"""
    root = _homebrew_root(instance_id, bot_id)
    if not root.exists():
        return []
    out = []
    for child in root.iterdir():
        if not child.is_dir():
            continue
        if not _safe_group_id(child.name):
            continue
        dbs = [p for p in child.iterdir() if p.is_file() and p.suffix.lower() == ".db"]
        out.append({
            "group_id": child.name,
            "db_count": len(dbs),
            "db_files": sorted(p.name for p in dbs),
        })
    return sorted(out, key=lambda x: x["group_id"])


def list_databases(instance_id: str, bot_id: str, group_id: str) -> List[Dict]:
    if not _safe_group_id(group_id):
        return []
    d = _group_dir(instance_id, bot_id, group_id)
    if not d.exists():
        return []
    out = []
    for p in d.iterdir():
        if not p.is_file() or p.suffix.lower() != ".db":
            continue
        try:
            conn = sqlite3.connect(str(p))
            try:
                row = conn.execute("SELECT COUNT(*) FROM data").fetchone()
                data_rows = row[0] if row else 0
            except sqlite3.DatabaseError:
                data_rows = 0
            try:
                row = conn.execute("SELECT COUNT(*) FROM redirect").fetchone()
                redirect_rows = row[0] if row else 0
            except sqlite3.DatabaseError:
                redirect_rows = 0
            conn.close()
            out.append({
                "name": p.name,
                "size": p.stat().st_size,
                "data_rows": data_rows,
                "redirect_rows": redirect_rows,
            })
        except sqlite3.DatabaseError:
            continue
    return sorted(out, key=lambda x: x["name"])


def create_database(instance_id: str, bot_id: str, group_id: str, name: str) -> Dict:
    """创建空白私设 db（schema 跟 QueryStore.create_empty_database 一致）。"""
    if not _safe_group_id(group_id):
        return {"ok": False, "message": "非法 group_id"}
    rel = _safe_db_filename(name)
    if rel is None:
        return {"ok": False, "message": "非法文件名（必须 .db 后缀，无路径分隔符）"}
    d = _group_dir(instance_id, bot_id, group_id)
    d.mkdir(parents=True, exist_ok=True)
    path = d / rel
    if path.exists():
        return {"ok": False, "message": "已存在同名 db"}
    try:
        conn = sqlite3.connect(str(path))
        try:
            data_cols = ", ".join(f'"{c}" TEXT DEFAULT (\'\')' for c in QUERY_DATA_FIELDS)
            redirect_cols = ", ".join(f'"{c}" TEXT DEFAULT (\'\')' for c in QUERY_REDIRECT_FIELDS)
            conn.execute(f"CREATE TABLE data ({data_cols})")
            conn.execute(f"CREATE TABLE redirect ({redirect_cols})")
            conn.execute('CREATE INDEX "idx_from" ON data ("来源")')
            conn.execute('CREATE INDEX "idx_catalogue" ON data ("分类")')
            conn.commit()
        finally:
            conn.close()
        return {"ok": True, "path": str(path)}
    except (OSError, sqlite3.DatabaseError) as e:
        return {"ok": False, "message": str(e)}


def delete_database(instance_id: str, bot_id: str, group_id: str, name: str) -> Dict:
    if not _safe_group_id(group_id):
        return {"ok": False, "message": "非法 group_id"}
    rel = _safe_db_filename(name)
    if rel is None:
        return {"ok": False, "message": "非法文件名"}
    path = _group_dir(instance_id, bot_id, group_id) / rel
    if not path.exists():
        return {"ok": False, "message": "db 不存在"}
    try:
        path.unlink()
        return {"ok": True}
    except OSError as e:
        return {"ok": False, "message": str(e)}


# ─── 条目 CRUD ───────────────────────────────────────────────────────────

def list_entries(instance_id: str, bot_id: str, group_id: str, db_name: str,
                 table: str = "data", offset: int = 0, limit: int = 200,
                 keyword: Optional[str] = None) -> Dict:
    if not _safe_group_id(group_id):
        return {"entries": [], "total": 0, "fields": []}
    rel = _safe_db_filename(db_name)
    if rel is None or table not in ("data", "redirect"):
        return {"entries": [], "total": 0, "fields": []}
    fields = QUERY_DATA_FIELDS if table == "data" else QUERY_REDIRECT_FIELDS
    path = _group_dir(instance_id, bot_id, group_id) / rel

    with _open_db(path) as conn:
        if conn is None:
            return {"entries": [], "total": 0, "fields": fields}
        where = ""
        params: List = []
        if keyword:
            clauses = [f'"{f}" LIKE ?' for f in fields]
            where = " WHERE " + " OR ".join(clauses)
            params = [f"%{keyword}%"] * len(fields)
        try:
            total = conn.execute(
                f'SELECT COUNT(*) FROM "{table}"{where}', params
            ).fetchone()[0]
            rows = conn.execute(
                f'SELECT rowid, * FROM "{table}"{where} ORDER BY rowid LIMIT ? OFFSET ?',
                params + [limit, offset],
            ).fetchall()
        except sqlite3.DatabaseError as e:
            return {"entries": [], "total": 0, "fields": fields, "error": str(e)}
        return {
            "entries": [dict(r) for r in rows],
            "total": total,
            "fields": fields,
        }


def upsert_entry(instance_id: str, bot_id: str, group_id: str, db_name: str,
                 table: str, rowid: Optional[int], values: Dict[str, str]) -> Dict:
    if not _safe_group_id(group_id) or table not in ("data", "redirect"):
        return {"ok": False, "message": "非法参数"}
    rel = _safe_db_filename(db_name)
    if rel is None:
        return {"ok": False, "message": "非法 db 名"}
    fields = QUERY_DATA_FIELDS if table == "data" else QUERY_REDIRECT_FIELDS
    clean = {f: str(values.get(f, "") or "") for f in fields}
    path = _group_dir(instance_id, bot_id, group_id) / rel

    with _open_db(path) as conn:
        if conn is None:
            return {"ok": False, "message": "db 不存在"}
        try:
            if rowid is None or rowid <= 0:
                cols = ", ".join(f'"{f}"' for f in fields)
                placeholders = ", ".join(["?"] * len(fields))
                cur = conn.execute(
                    f'INSERT INTO "{table}" ({cols}) VALUES ({placeholders})',
                    [clean[f] for f in fields],
                )
                conn.commit()
                return {"ok": True, "rowid": cur.lastrowid}
            else:
                set_clause = ", ".join(f'"{f}" = ?' for f in fields)
                cur = conn.execute(
                    f'UPDATE "{table}" SET {set_clause} WHERE rowid = ?',
                    [clean[f] for f in fields] + [rowid],
                )
                conn.commit()
                return {"ok": True, "rowid": rowid, "updated": cur.rowcount or 0}
        except sqlite3.DatabaseError as e:
            return {"ok": False, "message": str(e)}


def delete_entry(instance_id: str, bot_id: str, group_id: str, db_name: str,
                 table: str, rowid: int) -> Dict:
    if not _safe_group_id(group_id) or table not in ("data", "redirect"):
        return {"ok": False, "message": "非法参数"}
    rel = _safe_db_filename(db_name)
    if rel is None:
        return {"ok": False, "message": "非法 db 名"}
    path = _group_dir(instance_id, bot_id, group_id) / rel
    with _open_db(path) as conn:
        if conn is None:
            return {"ok": False, "message": "db 不存在"}
        try:
            cur = conn.execute(f'DELETE FROM "{table}" WHERE rowid = ?', (rowid,))
            conn.commit()
            return {"ok": True, "deleted": cur.rowcount or 0}
        except sqlite3.DatabaseError as e:
            return {"ok": False, "message": str(e)}


# ─── 上传 xlsx → db ──────────────────────────────────────────────────────

def upload_xlsx(instance_id: str, bot_id: str, group_id: str, db_name: str,
                xlsx_bytes: bytes) -> Dict:
    """把上传的 xlsx 内容转成 data 表条目追加到目标 db。

    xlsx 列：名称 | 英文 | 来源 | 分类 | 标签 | 内容（首行表头）
    """
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ImportError:
        return {"ok": False, "message": "openpyxl 未安装，无法解析 xlsx"}
    if not _safe_group_id(group_id):
        return {"ok": False, "message": "非法 group_id"}
    rel = _safe_db_filename(db_name)
    if rel is None:
        return {"ok": False, "message": "非法 db 名"}

    d = _group_dir(instance_id, bot_id, group_id)
    d.mkdir(parents=True, exist_ok=True)
    path = d / rel
    # 不存在则先创建 schema
    if not path.exists():
        create_result = create_database(instance_id, bot_id, group_id, db_name)
        if not create_result.get("ok"):
            return create_result

    # 解析 xlsx
    import io
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    except Exception as e:
        return {"ok": False, "message": f"xlsx 解析失败: {e}"}

    inserted = 0
    skipped = 0
    with _open_db(path) as conn:
        if conn is None:
            return {"ok": False, "message": "db 打开失败"}
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(c or "").strip() for c in rows[0]]
            # 列名映射到 QUERY_DATA_FIELDS 索引
            col_index: Dict[str, int] = {}
            for i, name in enumerate(header):
                if name in QUERY_DATA_FIELDS:
                    col_index[name] = i
            if "名称" not in col_index:
                skipped += len(rows) - 1
                continue
            cols = ", ".join(f'"{f}"' for f in QUERY_DATA_FIELDS)
            placeholders = ", ".join(["?"] * len(QUERY_DATA_FIELDS))
            for row in rows[1:]:
                if not row:
                    continue
                values = []
                for f in QUERY_DATA_FIELDS:
                    if f in col_index:
                        v = row[col_index[f]]
                        values.append(str(v) if v is not None else "")
                    else:
                        values.append("")
                if not values[0].strip():
                    skipped += 1
                    continue
                try:
                    conn.execute(
                        f'INSERT INTO data ({cols}) VALUES ({placeholders})',
                        values,
                    )
                    inserted += 1
                except sqlite3.DatabaseError:
                    skipped += 1
        conn.commit()
    return {"ok": True, "inserted": inserted, "skipped": skipped, "path": str(path)}


# ─── 群级宏（从 bot_data.db 的 group_macro 表读） ────────────────────────

def _bot_data_db(instance_id: str, bot_id: str) -> Path:
    return _bots_root(instance_id) / bot_id / "bot_data.db"


def list_group_macros(instance_id: str, bot_id: str, group_id: str) -> List[Dict]:
    path = _bot_data_db(instance_id, bot_id)
    if not path.exists():
        return []
    conn = sqlite3.connect(str(path))
    conn.row_factory = sqlite3.Row
    try:
        try:
            rows = conn.execute(
                "SELECT * FROM group_macro WHERE group_id = ? ORDER BY key",
                (group_id,),
            ).fetchall()
        except sqlite3.DatabaseError:
            return []
        return [dict(r) for r in rows]
    finally:
        conn.close()
